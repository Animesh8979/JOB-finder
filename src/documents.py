"""Render a tailored resume (dict) and cover letter (text) to DOCX and PDF files.

- DOCX is the nicely formatted, ATS-friendly file (via python-docx).
- PDF is produced with fpdf2 (no Microsoft Word required, so it always works).
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from . import config

# --- filenames ---------------------------------------------------------------
def slugify(text: str, maxlen: int = 50) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "_", (text or "").strip()).strip("_")
    return (s or "untitled")[:maxlen]


def output_stem(job: dict[str, Any], kind: str) -> Path:
    name = f"{slugify(job.get('company',''))}_{slugify(job.get('title',''))}_{kind}"
    return config.OUTPUTS_DIR / name


# --- DOCX --------------------------------------------------------------------
def hex_to_rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip("#")
    if len(hex_str) == 3:
        hex_str = "".join(c*2 for c in hex_str)
    r = int(hex_str[0:2], 16)
    g = int(hex_str[2:4], 16)
    b = int(hex_str[4:6], 16)
    return RGBColor(r, g, b)


def _base_doc(font_name: str = "Calibri", margin_mm: int = 18) -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(10.5)
    
    # Apply margin to all sections
    margin_in = margin_mm / 25.4  # convert to inches
    from docx.shared import Inches
    for section in doc.sections:
        section.top_margin = Inches(margin_in)
        section.bottom_margin = Inches(margin_in)
        section.left_margin = Inches(margin_in)
        section.right_margin = Inches(margin_in)
        
    return doc


def _heading(doc: Document, text: str, font_name: str = "Calibri", accent_rgb: RGBColor = RGBColor(0x1F, 0x4E, 0x79), template: str = "Corporate") -> None:
    p = doc.add_paragraph()
    
    if template == "Minimalist":
        run = p.add_run(text.title())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = font_name
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif template == "Modern":
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = font_name
        run.font.color.rgb = accent_rgb
    else:  # Corporate
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.name = font_name
        run.font.color.rgb = accent_rgb


def _edu_line(e: Any) -> str:
    if isinstance(e, dict):
        return " ".join(str(e.get(k, "")) for k in ("degree", "field", "school", "year")).strip()
    return str(e)


def save_resume_docx(resume: dict[str, Any], path: Path, style_config: dict[str, Any] | None = None) -> None:
    style = style_config or {}
    font_name = style.get("font", "Calibri")
    template = style.get("template", "Corporate")
    accent_hex = style.get("accent_color", "#1F4E79")
    accent_rgb = hex_to_rgb(accent_hex)
    margin_mm = style.get("margin", 18)

    doc = _base_doc(font_name, margin_mm)

    p = doc.add_paragraph()
    r = p.add_run(resume.get("name") or "Resume")
    r.bold = True
    r.font.name = font_name
    
    if template == "Modern":
        r.font.size = Pt(22)
        r.font.color.rgb = accent_rgb
    elif template == "Minimalist":
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor(0, 0, 0)
    else:  # Corporate
        r.font.size = Pt(20)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    if resume.get("contact"):
        c = doc.add_paragraph()
        run_c = c.add_run(resume["contact"])
        run_c.font.size = Pt(9.5)
        run_c.font.name = font_name
        if template == "Minimalist":
            run_c.italic = True

    if resume.get("summary"):
        _heading(doc, "Summary", font_name, accent_rgb, template)
        p_sum = doc.add_paragraph()
        run_sum = p_sum.add_run(resume["summary"])
        run_sum.font.name = font_name

    if resume.get("skills"):
        _heading(doc, "Skills", font_name, accent_rgb, template)
        p_sk = doc.add_paragraph()
        run_sk = p_sk.add_run(", ".join(map(str, resume["skills"])))
        run_sk.font.name = font_name

    if resume.get("experience"):
        _heading(doc, "Experience", font_name, accent_rgb, template)
        for job in resume["experience"]:
            line = doc.add_paragraph()
            run_t = line.add_run(str(job.get("title", "")).strip())
            run_t.bold = True
            run_t.font.name = font_name
            if job.get("company"):
                run_comp = line.add_run(f" — {job['company']}")
                run_comp.font.name = font_name
            meta = "  |  ".join(b for b in (job.get("location", ""), job.get("dates", "")) if b)
            if meta:
                m = doc.add_paragraph()
                mr = m.add_run(meta)
                mr.italic = True
                mr.font.size = Pt(9)
                mr.font.name = font_name
            for bullet in (job.get("bullets") or []):
                p_b = doc.add_paragraph(style="List Bullet")
                run_b = p_b.add_run(str(bullet))
                run_b.font.name = font_name

    if resume.get("projects"):
        _heading(doc, "Projects", font_name, accent_rgb, template)
        for pr in resume["projects"]:
            pp = doc.add_paragraph()
            run_pn = pp.add_run(str(pr.get("name", "")))
            run_pn.bold = True
            run_pn.font.name = font_name
            if pr.get("line"):
                run_pl = pp.add_run(f" — {pr['line']}")
                run_pl.font.name = font_name

    if resume.get("education"):
        _heading(doc, "Education", font_name, accent_rgb, template)
        for e in resume["education"]:
            p_e = doc.add_paragraph()
            run_e = p_e.add_run(_edu_line(e))
            run_e.font.name = font_name

    if resume.get("certifications"):
        _heading(doc, "Certifications", font_name, accent_rgb, template)
        p_cer = doc.add_paragraph()
        run_cer = p_cer.add_run(", ".join(map(str, resume["certifications"])))
        run_cer.font.name = font_name

    doc.save(str(path))


def save_cover_letter_docx(text: str, path: Path, name: str = "", contact: str = "", style_config: dict[str, Any] | None = None) -> None:
    style = style_config or {}
    font_name = style.get("font", "Calibri")
    margin_mm = style.get("margin", 18)
    
    doc = _base_doc(font_name, margin_mm)
    if name:
        h = doc.add_paragraph()
        r = h.add_run(name)
        r.bold = True
        r.font.size = Pt(14)
        r.font.name = font_name
    if contact:
        p_c = doc.add_paragraph()
        run_c = p_c.add_run(contact)
        run_c.font.size = Pt(9.5)
        run_c.font.name = font_name
    doc.add_paragraph("")
    for para in re.split(r"\n\s*\n", (text or "").strip()):
        p_p = doc.add_paragraph()
        run_p = p_p.add_run(para.strip())
        run_p.font.name = font_name
    doc.save(str(path))


# --- PDF (fpdf2; no Word needed) ---------------------------------------------
_REPL = {"–": "-", "—": "-", "•": "-", "“": '"', "”": '"', "‘": "'", "’": "'",
         "…": "...", " ": " ", "–": "-", "—": "-"}


def _latin1(s: str) -> str:
    for a, b in _REPL.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


def save_text_pdf(body: str, path: Path, title: str = "", contact: str = "", style_config: dict[str, Any] | None = None) -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    style = style_config or {}
    font_name = style.get("font", "Calibri")
    margin_mm = style.get("margin", 18)

    # Map standard fonts to fpdf core fonts (Helvetica, Times, Courier)
    fpdf_font = "Helvetica"
    if font_name in ("Georgia", "Times New Roman"):
        fpdf_font = "Times"

    def line(pdf: FPDF, text: str, h: float) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, h, _latin1(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=margin_mm)
    pdf.set_margins(margin_mm, margin_mm, margin_mm)
    pdf.add_page()
    if title:
        pdf.set_font(fpdf_font, "B", 16)
        line(pdf, title, 8)
    if contact:
        pdf.set_font(fpdf_font, "", 9)
        line(pdf, contact, 5)
        pdf.ln(2)
    pdf.set_font(fpdf_font, "", 11)
    for raw in (body or "").split("\n"):
        if raw.strip() == "":
            pdf.ln(3)
        else:
            line(pdf, raw, 5.5)
    pdf.output(str(path))


def resume_to_text(resume: dict[str, Any]) -> str:
    out: list[str] = []
    if resume.get("summary"):
        out += ["SUMMARY", resume["summary"], ""]
    if resume.get("skills"):
        out += ["SKILLS", ", ".join(map(str, resume["skills"])), ""]
    if resume.get("experience"):
        out.append("EXPERIENCE")
        for j in resume["experience"]:
            head = f"{j.get('title','')} - {j.get('company','')}".strip(" -")
            meta = "  |  ".join(b for b in (j.get("location", ""), j.get("dates", "")) if b)
            out.append(head + (f"  ({meta})" if meta else ""))
            for b in (j.get("bullets") or []):
                out.append(f"  - {b}")
            out.append("")
    if resume.get("projects"):
        out.append("PROJECTS")
        for p in resume["projects"]:
            out.append(f"{p.get('name','')} - {p.get('line','')}".strip(" -"))
        out.append("")
    if resume.get("education"):
        out.append("EDUCATION")
        out += [_edu_line(e) for e in resume["education"]]
        out.append("")
    if resume.get("certifications"):
        out += ["CERTIFICATIONS", ", ".join(map(str, resume["certifications"]))]
    return "\n".join(out)


# --- convenience -------------------------------------------------------------
def generate_documents(job: dict[str, Any], resume_dict: dict[str, Any], cover_text: str, style_config: dict[str, Any] | None = None) -> dict[str, str]:
    """Write all four files and return their paths."""
    name = resume_dict.get("name", "")
    contact = resume_dict.get("contact", "")
    rstem, cstem = output_stem(job, "resume"), output_stem(job, "cover_letter")

    paths: dict[str, str] = {}
    rdocx = rstem.with_suffix(".docx")
    save_resume_docx(resume_dict, rdocx, style_config)
    paths["resume_docx"] = str(rdocx)

    rpdf = rstem.with_suffix(".pdf")
    save_text_pdf(resume_to_text(resume_dict), rpdf, title=name, contact=contact, style_config=style_config)
    paths["resume_pdf"] = str(rpdf)

    cdocx = cstem.with_suffix(".docx")
    save_cover_letter_docx(cover_text, cdocx, name, contact, style_config)
    paths["cover_docx"] = str(cdocx)

    cpdf = cstem.with_suffix(".pdf")
    save_text_pdf(cover_text, cpdf, title=name, contact=contact, style_config=style_config)
    paths["cover_pdf"] = str(cpdf)
    return paths

