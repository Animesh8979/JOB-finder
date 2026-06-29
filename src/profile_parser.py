"""Turn an uploaded resume (PDF / DOCX / TXT) into a structured profile via the AI.

The resulting ``profile`` dict is the single source of truth used everywhere else:
- a compact text rendering of it is sent as the *cached* context for job scoring and
  document tailoring (so we only pay for it once across many calls);
- ``raw_text`` keeps the original resume text so tailoring can ground itself in real
  wording and never invents experience.
"""
from __future__ import annotations

import io
from typing import Any

from . import llm

# The shape we ask the model to return. Kept flat and predictable for the UI.
PROFILE_KEYS = (
    "name", "email", "phone", "location", "headline", "summary", "skills",
    "years_experience", "experience", "education", "certifications", "projects",
    "links", "target_titles", "ats_keyword_aliases"
)

_PARSE_SYSTEM = (
    "You are an expert resume parser. Extract the candidate's information faithfully. "
    "Do not invent or embellish anything. Extract hard technical skills and explicitly normalize them "
    "into an 'ats_keyword_aliases' list (e.g., mapping 'ReactJS' to ['React.js', 'React']). "
    "If a field is unknown, use an empty string or empty list."
)


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Pull plain text out of a PDF, DOCX, or TXT resume."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _pdf_text(file_bytes)
    if name.endswith(".docx"):
        return _docx_text(file_bytes)
    if name.endswith(".txt"):
        return file_bytes.decode("utf-8", "ignore")
    raise ValueError("Please upload a PDF, Word (.docx), or .txt resume.")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _docx_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def parse_resume(text: str) -> dict[str, Any]:
    """Ask the AI to structure the resume text into a profile dict."""
    if not text.strip():
        raise ValueError("Couldn't read any text from that file. Is it a scanned image PDF?")

    prompt = (
        "Extract this resume into JSON with EXACTLY these keys:\n"
        "name (str), email (str), phone (str), location (str), headline (str, a short "
        "professional title line), summary (str, 2-3 sentences), skills (list of str), "
        "years_experience (number), "
        "experience (list of objects: title, company, location, start, end, bullets[list of str]), "
        "education (list of objects: degree, field, school, year), "
        "certifications (list of str), "
        "projects (list of objects: name, description, tech[list of str]), "
        "links (object: linkedin, github, portfolio), "
        "target_titles (list of str — infer 2-4 job titles this person fits), "
        "ats_keyword_aliases (list of str — normalized ATS keywords and variations).\n\n"
        "RESUME TEXT:\n\"\"\"\n" + text + "\n\"\"\""
    )
    data = llm.generate_json(prompt, system=_PARSE_SYSTEM, max_tokens=3500)
    if not isinstance(data, dict):
        raise llm.LLMError("The AI returned an unexpected format. Please try again.")
    # Ensure all keys exist so the UI never KeyErrors.
    for key in PROFILE_KEYS:
        data.setdefault(key, "" if key in ("name", "email", "phone", "location", "headline", "summary") else [])
    if not isinstance(data.get("links"), dict):
        data["links"] = {"linkedin": "", "github": "", "portfolio": ""}
    data["raw_text"] = text
    return data


def profile_context(profile: dict[str, Any]) -> str:
    """Render the profile to a compact, stable text block used as cached AI context."""
    if not profile:
        return ""
    lines: list[str] = ["=== CANDIDATE PROFILE ==="]
    if profile.get("headline"):
        lines.append(f"Headline: {profile['headline']}")
    if profile.get("location"):
        lines.append(f"Location: {profile['location']}")
    if profile.get("years_experience"):
        lines.append(f"Years of experience: {profile['years_experience']}")
    if profile.get("summary"):
        lines.append(f"Summary: {profile['summary']}")
    if profile.get("skills"):
        lines.append("Skills: " + ", ".join(map(str, profile["skills"])))

    if profile.get("experience"):
        lines.append("\nExperience:")
        for job in profile["experience"]:
            header = f"- {job.get('title','')} @ {job.get('company','')}".rstrip(" @")
            span = " ".join(x for x in (job.get("start", ""), "–", job.get("end", "")) if x).strip(" –")
            if span:
                header += f" ({span})"
            lines.append(header)
            for bullet in (job.get("bullets") or [])[:6]:
                lines.append(f"    • {bullet}")

    if profile.get("education"):
        lines.append("\nEducation:")
        for ed in profile["education"]:
            lines.append(
                f"- {ed.get('degree','')} {ed.get('field','')}, {ed.get('school','')} {ed.get('year','')}".strip()
            )
    if profile.get("certifications"):
        lines.append("Certifications: " + ", ".join(map(str, profile["certifications"])))
    if profile.get("projects"):
        lines.append("\nProjects:")
        for pr in profile["projects"]:
            tech = ", ".join(pr.get("tech") or [])
            lines.append(f"- {pr.get('name','')}: {pr.get('description','')}" + (f" [{tech}]" if tech else ""))

    # The original text grounds the AI and prevents fabrication during tailoring.
    if profile.get("raw_text"):
        lines.append("\n=== ORIGINAL RESUME TEXT (authoritative; do not contradict) ===")
        lines.append(profile["raw_text"])
    return "\n".join(lines)
